#!/usr/bin/env python3
"""Generate vanilla/doom-fire.xlsx — Doom fire in plain Excel.

No macros, no LET, no dynamic arrays.  Formulas + conditional formatting only.
Upstream provenance: playground/1st-implementation-with-tables/fire.js @854c39f

Layout (SPEC-003):
  Fire  (first sheet)  A1:BH40  — 60 cols x 40 rows, row 40 = =Source
  Decay (second sheet) A1:BH39  — =RANDBETWEEN(0,2) per cell (fire.js L43)

Propagation (D-04 / SPEC-007, pull form of fire.js L35-49):
  For r in 1..39, c in 1..60, d = Decay!(r,c):
    Fire!(r,c) = MAX(0, Fire!(r+1, c+d) - d), or 0 when c+d > 60.
  Right-edge overflow (c+d > 60) reads helper columns BI/BJ which hold 0.

Rendering (D-06 / SPEC-010/011):
  37 cellIs/equal CF rules, one per intensity 0..36, exact RGB from fire.js L5.
  Number format ;;;  gridlines off  uniform square-ish cells.
"""
import argparse
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

ROWS, COLS = 40, 60
SOURCE_DEFAULT = 36
SOURCE_CELL = "A42"           # outside the grid (row 42 > 40)
COL_WIDTH = 2                # Excel char units → px = width*7+5 = 19
ROW_HEIGHT = 15               # points → px = height*96/72 = 20, ratio 0.95

REPO = Path(__file__).resolve().parents[3]
FIREJS = REPO / "playground" / "1st-implementation-with-tables" / "fire.js"


def upstream_palette():
    """37 (r, g, b) tuples parsed from fire.js L5 @854c39f."""
    line = FIREJS.read_text().splitlines()[4]
    return [
        (int(r), int(g), int(b))
        for r, g, b in re.findall(r'"r":(\d+),"g":(\d+),"b":(\d+)', line)
    ]


def build(out_path: str) -> None:
    palette = upstream_palette()
    assert len(palette) == 37

    wb = Workbook()

    # ── Fire sheet (first) ────────────────────────────────────────────
    fire = wb.active
    fire.title = "Fire"

    # Source input cell (outside grid) + whole-number validation 0..36
    fire[SOURCE_CELL] = SOURCE_DEFAULT
    dv = DataValidation(
        type="whole", operator="between", formula1="0", formula2="36"
    )
    dv.add(SOURCE_CELL)
    fire.add_data_validation(dv)

    # Defined name Source → Fire!$A$42  (SPEC-003)
    wb.defined_names["Source"] = DefinedName(
        "Source", attr_text=f"Fire!${SOURCE_CELL[0]}${SOURCE_CELL[1:]}"
    )

    # Row 40 (source row): every cell = =Source  (SPEC-004)
    for c in range(1, COLS + 1):
        fire.cell(ROWS, c).value = "=Source"

    # Rows 1..39: propagation formula  (SPEC-007)
    #   Fire!(r,c) = MAX(0, INDEX(Fire!A(r+1):BJ(r+1), 1, c + Decay!(r,c))
    #                     - Decay!(r,c))
    # Helper cols BI(61)/BJ(62) = 0 absorb right-edge overflow (c+d > 60).
    for r in range(1, ROWS):
        next_row = r + 1
        row_range = f"Fire!A{next_row}:BJ{next_row}"
        for c in range(1, COLS + 1):
            decay_ref = f"Decay!{get_column_letter(c)}{r}"
            fire.cell(r, c).value = (
                f"=MAX(0,INDEX({row_range},1,{c}+{decay_ref})-{decay_ref})"
            )

    # Helper columns BI (61) and BJ (62): zero buffer for right edge
    for r in range(1, ROWS + 1):
        fire.cell(r, 61).value = 0
        fire.cell(r, 62).value = 0

    # Number format ;;; for grid cells  (SPEC-011)
    for r in range(1, ROWS + 1):
        for c in range(1, COLS + 1):
            fire.cell(r, c).number_format = ";;;"

    # Gridlines off  (SPEC-011)
    fire.sheet_view.showGridLines = False

    # Uniform square-ish column widths / row heights  (SPEC-011)
    for c in range(1, COLS + 1):
        fire.column_dimensions[get_column_letter(c)].width = COL_WIDTH
    for r in range(1, ROWS + 1):
        fire.row_dimensions[r].height = ROW_HEIGHT

    # 37 cellIs/equal CF rules, one per intensity 0..36  (SPEC-010)
    cf_range = f"A1:{get_column_letter(COLS)}{ROWS}"
    for k, (r, g, b) in enumerate(palette):
        rgb = f"FF{r:02X}{g:02X}{b:02X}"
        fill = PatternFill(start_color=rgb, end_color=rgb, fill_type="solid")
        fire.conditional_formatting.add(
            cf_range,
            CellIsRule(operator="equal", formula=[str(k)], fill=fill),
        )

    # ── Decay sheet ───────────────────────────────────────────────────
    decay = wb.create_sheet("Decay")
    for r in range(1, ROWS):                       # rows 1..39
        for c in range(1, COLS + 1):               # cols A..BH
            decay.cell(r, c).value = "=RANDBETWEEN(0,2)"
    # row 40 of Decay is intentionally empty

    # ── Save ───────────────────────────────────────────────────────────
    p = Path(out_path)
    p.parent.mkdir(parents=True, exist_ok=True)
    wb.save(p)


def main():
    parser = argparse.ArgumentParser(
        description="Generate vanilla Doom fire .xlsx (no macros)."
    )
    parser.add_argument("--out", required=True, help="Output .xlsx path")
    args = parser.parse_args()
    build(args.out)
    print(f"Wrote {args.out}")


if __name__ == "__main__":
    main()

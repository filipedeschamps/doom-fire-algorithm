"""Hermetic guards: harness validity + upstream provenance. No artifact needed."""
import openpyxl
import pytest

from conftest import ROWS, COLS, upstream_palette, FIREJS
from support import recalc as rc
from support.invariant import propagation_violations

TRACES = ["SPEC-007", "SPEC-008", "SPEC-010", "SPEC-003"]


def test_guard_upstream_constants_pinned():
    lines = FIREJS.read_text().splitlines()
    assert lines[1].strip() == "let fireWidth = 60"
    assert lines[2].strip() == "let fireHeight = 40"
    assert (ROWS, COLS) == (40, 60)
    p = upstream_palette()
    assert len(p) == 37 and p[0] == (7, 7, 7) and p[36] == (255, 255, 255)


def test_guard_invariant_checker_detects_errors():
    w = 5
    fire = [[0] * w for _ in range(3)]
    fire[2] = [10] * w
    decay = [[0] * w, [1] * w]
    # row1: c+d>4 at last col -> 0; others 10-1 = 9
    fire[1] = [9, 9, 9, 9, 0]
    # row0: d=0 -> copy of row1
    fire[0] = list(fire[1])
    assert propagation_violations(fire, decay) == []
    fire[0][2] = 8
    assert propagation_violations(fire, decay) == [(0, 2, 9, 8)]
    fire[0][2] = 9
    fire[1][4] = 9  # right edge must read 0
    assert propagation_violations(fire, decay)


def test_guard_libreoffice_recalc_harness(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 7
    ws["B1"] = "=A1*2"
    for i in range(1, 41):
        ws.cell(i, 3).value = "=RANDBETWEEN(0,2)"
    from openpyxl.workbook.defined_name import DefinedName
    wb.defined_names["Source"] = DefinedName("Source", attr_text="Sheet!$A$1")
    ws.title = "Sheet"
    p = tmp_path / "t.xlsx"
    wb.save(p)
    a = rc.recalc(p, source=21)["Sheet"]
    b = rc.recalc(p, source=21)["Sheet"]
    assert a["A1"].value == 21 and a["B1"].value == 42
    col = lambda s: [s.cell(i, 3).value for i in range(1, 41)]
    assert set(col(a)) <= {0, 1, 2} and col(a) != col(b)

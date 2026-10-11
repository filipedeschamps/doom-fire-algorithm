"""SPEC-006 Excel 2010+/LibreOffice compatibility."""
import re
import zipfile

import openpyxl
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from conftest import COLS, ROWS

TRACES = ["SPEC-006"]

FORBIDDEN = [
    "LET", "LAMBDA", "FILTER", "SORT", "SORTBY", "UNIQUE", "SEQUENCE",
    "RANDARRAY", "XLOOKUP", "XMATCH", "MAP", "REDUCE", "SCAN", "BYROW",
    "BYCOL", "MAKEARRAY", "TAKE", "DROP", "CHOOSECOLS", "CHOOSEROWS",
    "HSTACK", "VSTACK", "TOCOL", "TOROW", "WRAPCOLS", "WRAPROWS",
    "TEXTSPLIT", "TEXTBEFORE", "TEXTAFTER", "IFS", "SWITCH", "MAXIFS",
    "MINIFS", "CONCAT", "TEXTJOIN",
]


def _formulas(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None and (
                    isinstance(c.value, (ArrayFormula, DataTableFormula))
                    or str(c.value).startswith("=")
                ):
                    yield ws.title, c.coordinate, c.value


def test_spec006_no_modern_functions_or_arrays(wb):
    pat = re.compile(r"\b(" + "|".join(FORBIDDEN) + r")\s*\(", re.I)
    n = 0
    for sheet, coord, v in _formulas(wb):
        n += 1
        assert not isinstance(v, (ArrayFormula, DataTableFormula)), (sheet, coord)
        s = str(v)
        assert "_xlfn" not in s.lower() and "_xlws" not in s.lower(), (sheet, coord)
        assert "#" not in s, (sheet, coord, "spill operator")
        assert not pat.search(s), (sheet, coord, s)
    assert n >= ROWS * COLS + (ROWS - 1) * COLS


def test_spec006_no_iterative_calc(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path)
    calc = wb.calculation
    assert not calc.iterate


def test_spec006_no_formula_errors_after_recalc(recalcs):
    book = recalcs(36)
    for name in ("Fire", "Decay"):
        for row in book[name].iter_rows():
            for c in row:
                v = c.value
                assert not (isinstance(v, str) and v.startswith(("#", "Err:"))), (
                    name, c.coordinate, v)


def test_spec006_no_external_links(xlsx_path):
    with zipfile.ZipFile(xlsx_path) as z:
        assert not [n for n in z.namelist() if "externalLink" in n]

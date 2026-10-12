"""SPEC-003 grid/layout and SPEC-004 source row."""
import pytest
from openpyxl.utils import get_column_letter as L

from conftest import COLS, ROWS, fire_decay
from support import recalc as rc
from support.invariant import propagation_violations  # noqa: F401

TRACES = ["SPEC-003", "SPEC-004"]


def test_spec003_sheet_names(wb):
    assert "Fire" in wb.sheetnames and "Decay" in wb.sheetnames
    assert wb.sheetnames[0] == "Fire"


def test_spec003_fire_grid_is_formulas_a1_bh40(wb):
    ws = wb["Fire"]
    assert L(COLS) == "BH"
    for r in range(1, ROWS + 1):
        for c in range(1, COLS + 1):
            v = ws.cell(r, c).value
            assert isinstance(v, str) and v.startswith("="), (r, c, v)


def test_spec003_source_named_single_input_cell(wb):
    sheet, ref = rc.source_cell(wb)
    assert sheet == "Fire"
    cell = wb[sheet][ref]
    assert cell.value == 36
    # outside the 60x40 fire grid
    assert not (1 <= cell.row <= ROWS and 1 <= cell.column <= COLS)


def test_spec003_source_validation_whole_0_36(wb):
    sheet, ref = rc.source_cell(wb)
    ws = wb[sheet]
    hits = [
        dv
        for dv in ws.data_validations.dataValidation
        if ref in str(dv.sqref).replace("$", "").split()
    ]
    assert hits, "Source needs a data validation"
    dv = hits[0]
    assert dv.type == "whole"
    assert (dv.operator or "between") == "between"
    assert str(dv.formula1) == "0" and str(dv.formula2) == "36"


def test_spec003_decay_sheet_extent(wb):
    ws = wb["Decay"]
    assert ws.max_column == COLS and ws.max_row == ROWS - 1  # before any access
    for r in range(1, ROWS):
        for c in range(1, COLS + 1):
            assert ws.cell(r, c).value is not None, (r, c)
    assert all(ws.cell(ROWS, c).value is None for c in range(1, COLS + 1))


def test_spec004_bottom_row_formula_is_source(wb):
    ws = wb["Fire"]
    for c in range(1, COLS + 1):
        assert ws.cell(ROWS, c).value.replace(" ", "").upper() == "=SOURCE"


@pytest.mark.parametrize("src", [36, 0, 17])
def test_spec004_bottom_row_equals_source_after_recalc(recalcs, src):
    fire, _ = fire_decay(recalcs(src))
    assert fire[ROWS - 1] == [src] * COLS


def test_spec004_destroy_source_gives_all_zero(recalcs):
    fire, _ = fire_decay(recalcs(0))
    assert all(v == 0 for row in fire for v in row)


def test_spec004_create_source_lights_fire(recalcs):
    fire, _ = fire_decay(recalcs(36))
    assert fire[ROWS - 1] == [36] * COLS
    assert any(v > 0 for v in fire[ROWS - 2])
    assert max(fire[ROWS - 2]) >= 34  # one step loses at most 2

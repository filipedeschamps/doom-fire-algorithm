"""Recalculate a workbook with headless LibreOffice and read cached values.

No GUI Excel needed. The input is first round-tripped through openpyxl so no
stale cached values survive (forces a genuine recalculation of every formula,
volatile or not).
"""
import os
import shutil
import subprocess
import tempfile
from pathlib import Path

import openpyxl

SOFFICE = os.environ.get(
    "SOFFICE", "/Applications/LibreOffice.app/Contents/MacOS/soffice"
)


def recalc(xlsx_path, source=None, timeout=180):
    """Return a data_only openpyxl Workbook recalculated by LibreOffice.

    source: if not None, the defined name `Source` cell is set to this value
    before recalculation.
    """
    tmp = Path(tempfile.mkdtemp(prefix="doomfire-"))
    try:
        wb = openpyxl.load_workbook(xlsx_path)
        if source is not None:
            set_source(wb, source)
        src = tmp / "in" / "book.xlsx"
        src.parent.mkdir()
        wb.save(src)
        out = tmp / "out"
        out.mkdir()
        profile = tmp / "profile"
        cmd = [
            SOFFICE,
            f"-env:UserInstallation=file://{profile}",
            "--headless",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(out),
            str(src),
        ]
        subprocess.run(cmd, check=True, capture_output=True, timeout=timeout)
        return openpyxl.load_workbook(out / "book.xlsx", data_only=True)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def source_cell(wb):
    """Resolve defined name `Source` to (sheet title, coordinate)."""
    dn = wb.defined_names["Source"]
    dests = list(dn.destinations)
    assert len(dests) == 1, "Source must name exactly one cell"
    sheet, ref = dests[0]
    ref = ref.replace("$", "")
    assert ":" not in ref, "Source must be a single cell"
    return sheet, ref


def set_source(wb, value):
    sheet, ref = source_cell(wb)
    wb[sheet][ref] = value


def grid(ws, rows, cols, top=1, left=1):
    """rows x cols list-of-lists of cell values."""
    return [
        [ws.cell(row=top + r, column=left + c).value for c in range(cols)]
        for r in range(rows)
    ]
